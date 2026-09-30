# -*- coding: utf-8 -*-
"""单元计算查漏迷你纸模板（1 页）· 天蓝主色。"""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "printables" / "templates"
STEM = "template-unit-calc-check"
sys.path.insert(0, str(Path(__file__).resolve().parent))
from theme import BRAND_LINE, RGB_SKY_DEEP, RGB_SKY_DARK  # noqa: E402


def set_run_font(run, name="宋体", size=11, bold=False, color=None):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def add_para(doc, text, *, size=11, bold=False, align=None, space_after=4, color=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(space_after)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return p


def build() -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    for m in ("left_margin", "right_margin"):
        setattr(sec, m, Cm(1.6))
    sec.top_margin = Cm(1.2)
    sec.bottom_margin = Cm(1.2)

    add_para(doc, BRAND_LINE, size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
             color=RGB_SKY_DEEP, space_after=2)
    add_para(doc, "单元计算查漏 · 迷你纸（模板）", size=13, bold=True,
             align=WD_ALIGN_PARAGRAPH.CENTER, color=RGB_SKY_DARK, space_after=2)
    add_para(doc, "建议 8–12 分钟　|　只练计算查漏　|　思维题请另用短思维纸",
             size=9, align=WD_ALIGN_PARAGRAPH.CENTER, color=RGB_SKY_DEEP, space_after=6)

    add_para(doc, "单元：第____单元「____________________」　姓名：________　日期：________",
             size=10, space_after=8)

    add_para(doc, "一、口算 / 简便（约 3 分钟，不限时速写）", size=12, bold=True,
             color=RGB_SKY_DARK, space_after=4)
    for line in [
        "1. ______＋______＝______　　2. ______－______＝______　　3. ______×______＝______",
        "4. ______÷______＝______　　5. ______×______＝______　　6. ______÷______＝______",
        "7. ______　　8. ______　　9. ______　　10. ______",
    ]:
        add_para(doc, line, size=11, space_after=4)

    add_para(doc, "二、竖式（少量，练熟练与工整 · 约 4 分钟）", size=12, bold=True,
             color=RGB_SKY_DARK, space_after=4)
    add_para(doc, "列竖式计算。带 ※ 的要验算。", size=10, space_after=4)
    add_para(doc, "1. ________　　　2. ________　　　※3. ________", size=11, space_after=20)
    add_para(doc, "（下方自行留竖式书写空位）", size=9, color=RGB_SKY_DEEP, space_after=10)

    add_para(doc, "三、错因勾选（做完批改后勾）", size=12, bold=True,
             color=RGB_SKY_DARK, space_after=4)
    add_para(doc, "□ 计算错（进退位 / 口诀 / 数位对齐）　　□ 概念错（意义 / 单位 / 近似）",
             size=10, space_after=3)
    add_para(doc, "□ 审题错　　□ 书写潦草看错　　□ 其他：______________", size=10, space_after=8)

    add_para(doc, "四、订正区（计算错 / 概念错分两栏更好）", size=12, bold=True,
             color=RGB_SKY_DARK, space_after=4)
    add_para(doc, "【计算错重练】", size=10, bold=True, space_after=2)
    add_para(doc, "1. __________＝______　　2. __________＝______　　3. __________＝______",
             size=11, space_after=6)
    add_para(doc, "【概念错重练】用一句话写出正确概念：", size=10, bold=True, space_after=2)
    add_para(doc, "________________________________________________________________", size=11, space_after=2)
    add_para(doc, "________________________________________________________________", size=11, space_after=8)

    add_para(doc, f"{BRAND_LINE}　·　模板页，放入真实题目后另存到 printables/",
             size=8, align=WD_ALIGN_PARAGRAPH.CENTER, color=RGB_SKY_DEEP)
    return doc


def docx_to_pdf(docx_path: Path, pdf_path: Path) -> None:
    import win32com.client
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    try:
        d = word.Documents.Open(str(docx_path.resolve()))
        d.SaveAs(str(pdf_path.resolve()), FileFormat=17)
        d.Close(False)
    finally:
        word.Quit()


def main() -> int:
    OUT.mkdir(parents=True, exist_ok=True)
    docx_path = OUT / f"{STEM}.docx"
    pdf_path = OUT / f"{STEM}.pdf"
    build().save(str(docx_path))
    print(f"Wrote {docx_path}")
    try:
        docx_to_pdf(docx_path, pdf_path)
        print(f"Wrote {pdf_path}")
    except Exception as e:
        print(f"PDF skipped: {e}", file=sys.stderr)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
