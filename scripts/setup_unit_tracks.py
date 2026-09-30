# -*- coding: utf-8 -*-
"""按「单元 × 五条线」整理目录，并生成带时长的练习纸（Word+PDF）。"""
from __future__ import annotations

import shutil
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(Path(__file__).resolve().parent))
from theme import BRAND_LINE, RGB_SKY_DEEP, RGB_SKY_DARK  # noqa: E402

UNITS = [
    {
        "id": "u01",
        "name": "大数的认识",
        "exam_time": "40 分钟",
        "focus": [
            "数位 / 数级；读写大数（中间、末尾的 0）",
            "改写与省略（万、亿）；四舍五入求近似数",
            "比较大小；算盘 / 计数器 / 数线",
        ],
        "daily": "读数写数、数位意义、整万整亿口算/改写",
        "key": "近似数口诀、改写省略易错点、数线与算盘表示",
        "app": "情境中的大数读写、近似、比较（如距离、人口）",
        "source_glob": "U1-*.png",
        "printable_prefix": "U1_",
    },
    {
        "id": "u02",
        "name": "角的度量",
        "exam_time": "40 分钟",
        "focus": [
            "锐角 / 直角 / 钝角 / 平角 / 周角",
            "量角器内外圈；三角尺拼角",
            "角的和差；生活中的角",
        ],
        "daily": "估角、常见角度数、三角尺固定角",
        "key": "内外圈读错、拼角、对折圆/半圆求角",
        "app": "量角画角、转杆/椅背/滑雪道等情境",
        "source_glob": "U2-*.png",
        "printable_prefix": "U2_",
    },
    {
        "id": "u03",
        "name": "三位数乘两位数",
        "exam_time": "40 分钟",
        "focus": [
            "积的变化规律；估算（够不够）",
            "竖式部分积意义（×10 / ×几十）",
            "速度×时间＝路程等乘算应用",
        ],
        "daily": "两位数×一位数、整十乘、口算积",
        "key": "部分积错位、估算方法选择、积的变化",
        "app": "路程、面积增减、连乘情境",
        "source_glob": "U3.png",
        "printable_prefix": "U3_",
    },
    {
        "id": "u04",
        "name": "数量关系",
        "exam_time": "40 分钟",
        "focus": [
            "总价＝单价×数量",
            "路程＝速度×时间",
            "总数与部分；已知求未知量",
        ],
        "daily": "三类关系式填空与口算",
        "key": "问的是哪个量、单位换算后再比快慢",
        "app": "购物、行程、交通方式判断",
        "source_glob": "U4.png",
        "printable_prefix": "U4_",
    },
    {
        "id": "u05",
        "name": "平行四边形和梯形",
        "exam_time": "40 分钟",
        "focus": [
            "平行与垂直；点到直线的距离",
            "梯形 / 平行四边形判定；特殊四边形",
            "周长；过点作垂线 / 平行线",
        ],
        "daily": "辨平行垂直、量长度、简单周长",
        "key": "一组平行 vs 两组平行；等腰/直角梯形",
        "app": "作图与图形遮挡推理",
        "source_glob": "U5.png",
        "printable_prefix": "U5_",
    },
]

TRACKS = [
    ("01-日常计算", "约 15 分钟", "提速与熟练：口算 / 基础计算，题量适中"),
    ("02-重难点题型", "约 15 分钟", "本单元易错点、概念题、典型题型专练"),
    ("03-应用题", "约 15 分钟", "短应用 / 情境题；与口算分开，不混成大卷"),
    ("04-考前测试", "30–40 分钟", "原卷入库 + 延展卷（难度/题型扩展）"),
    ("05-错题库", "按错题量", "计算错 / 概念错分栏重练"),
]


def set_run_font(run, name="宋体", size=11, bold=False, color=None):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def add_para(doc, text, *, size=11, bold=False, align=None, space_after=4, color=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(space_after)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return p


def docx_to_pdf(docx_path: Path, pdf_path: Path) -> None:
    import win32com.client

    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    try:
        d = word.Documents.Open(str(docx_path.resolve()))
        d.SaveAs(str(pdf_path.resolve()), FileFormat=17)
        d.Close(False)
    finally:
        word.Quit()


def header(doc, unit_name: str, track: str, minutes: str, subtitle: str):
    add_para(doc, BRAND_LINE, size=15, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, color=RGB_SKY_DEEP)
    add_para(doc, f"四上 · {unit_name} · {track}", size=13, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, color=RGB_SKY_DARK)
    add_para(doc, f"限时 {minutes}　|　{subtitle}　|　番茄钟自计时", size=9, align=WD_ALIGN_PARAGRAPH.CENTER, color=RGB_SKY_DEEP)
    add_para(doc, "姓名：________　日期：________　完成：____　正确：____", size=10, space_after=8)


def build_daily(unit: dict) -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for m in ("left_margin", "right_margin"):
        setattr(sec, m, Cm(1.6))
    sec.top_margin = sec.bottom_margin = Cm(1.3)
    header(doc, unit["name"], "日常计算练习", "约 15 分钟", unit["daily"])
    add_para(doc, "一、口算（约 8 分钟，做完几题写几题）", size=12, bold=True, color=RGB_SKY_DARK)
    for i in range(1, 17):
        add_para(doc, f"{i}. __________ ＝ ______", size=11, space_after=3)
    add_para(doc, "二、少量竖式（约 5 分钟，练工整）", size=12, bold=True, color=RGB_SKY_DARK, space_after=4)
    add_para(doc, "1. ________　　2. ________　　3. ________（下方列竖式）", size=11, space_after=16)
    add_para(doc, "订正区：______________________________________________", size=10, color=RGB_SKY_DARK)
    add_para(doc, "（题目按教材与本单元试卷重难点填入；本页为时长版式模板。）", size=8, color=RGB_SKY_DEEP)
    return doc


def build_key(unit: dict) -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for m in ("left_margin", "right_margin"):
        setattr(sec, m, Cm(1.6))
    sec.top_margin = sec.bottom_margin = Cm(1.3)
    header(doc, unit["name"], "重难点题型练习", "约 15 分钟", unit["key"])
    add_para(doc, "本单元重难点（来自教材 + 单元卷）：", size=11, bold=True, color=RGB_SKY_DARK)
    for f in unit["focus"]:
        add_para(doc, f"· {f}", size=10, space_after=2)
    add_para(doc, "", size=8, space_after=6)
    add_para(doc, "一、概念 / 选择 / 判断（约 7 分钟）", size=12, bold=True, color=RGB_SKY_DARK)
    for i in range(1, 7):
        add_para(doc, f"{i}. ______________________________________________", size=11, space_after=4)
    add_para(doc, "二、典型题（约 8 分钟）", size=12, bold=True, color=RGB_SKY_DARK)
    for i in range(1, 5):
        add_para(doc, f"{i}. ______________________________________________", size=11, space_after=2)
        add_para(doc, "答：__________________________________________", size=11, space_after=6)
    return doc


def build_app(unit: dict) -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for m in ("left_margin", "right_margin"):
        setattr(sec, m, Cm(1.6))
    sec.top_margin = sec.bottom_margin = Cm(1.3)
    header(doc, unit["name"], "应用题练习", "约 15 分钟", unit["app"])
    add_para(doc, "说明：短应用纸，不与口算混印。先写数量关系再算。", size=9, color=RGB_SKY_DEEP, space_after=8)
    for i in range(1, 5):
        add_para(doc, f"{i}.（情境题，约 3–4 分钟）", size=11, bold=True, color=RGB_SKY_DARK)
        add_para(doc, "题目：______________________________________________", size=11, space_after=2)
        add_para(doc, "____________________________________________________", size=11, space_after=2)
        add_para(doc, "数量关系：____________________　　列式：____________", size=11, space_after=2)
        add_para(doc, "答：____________________", size=11, space_after=8)
    return doc


def build_wrong(unit: dict) -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for m in ("left_margin", "right_margin"):
        setattr(sec, m, Cm(1.6))
    sec.top_margin = sec.bottom_margin = Cm(1.3)
    header(doc, unit["name"], "错题重练纸", "按错题量", "计算错 / 概念错分开")
    add_para(doc, "来源：□日常计算 □重难点 □应用题 □考前测试　日期：________", size=10, space_after=8)
    add_para(doc, "【计算错】重练（对齐、进退位、口诀、数位）", size=12, bold=True, color=RGB_SKY_DARK)
    for i in range(1, 5):
        add_para(doc, f"{i}. 原错：__________　重做：__________ ＝ ______", size=11, space_after=4)
    add_para(doc, "【概念错】用一句话写对", size=12, bold=True, color=RGB_SKY_DARK, space_after=4)
    for i in range(1, 4):
        add_para(doc, f"{i}. 错因：__________　正解：______________________________", size=11, space_after=4)
    return doc


def ensure_dirs(unit: dict) -> Path:
    base = ROOT / "units" / f"{unit['id']}-{unit['name']}"
    for track, _, _ in TRACKS:
        (base / track).mkdir(parents=True, exist_ok=True)
    (base / "04-考前测试" / "source").mkdir(parents=True, exist_ok=True)
    (base / "04-考前测试" / "original").mkdir(parents=True, exist_ok=True)
    (base / "04-考前测试" / "extended").mkdir(parents=True, exist_ok=True)
    for track, _, _ in TRACKS:
        (ROOT / "printables" / unit["id"] / track).mkdir(parents=True, exist_ok=True)
    return base


def write_unit_readme(unit: dict, base: Path) -> None:
    lines = [
        f"# 四上 · {unit['name']}",
        "",
        f"**考前测试时长**：{unit['exam_time']}（固定 30–40 分钟）  ",
        "**日常计算 / 重难点 / 应用题**：每次约 **15 分钟**  ",
        "**主色**：天蓝",
        "",
        "## 五条线",
        "",
        "| 线 | 时长 | 说明 |",
        "| --- | --- | --- |",
    ]
    for track, mins, desc in TRACKS:
        lines.append(f"| {track} | {mins} | {desc} |")
    lines += [
        "",
        "## 本单元重难点（教材 PDF + 单元卷 JPG）",
        "",
    ]
    for f in unit["focus"]:
        lines.append(f"- {f}")
    lines += [
        "",
        "## 文件位置",
        "",
        f"- 可打印：`../../printables/{unit['id']}/`",
        f"- 考前原卷扫描：`04-考前测试/source/`",
        f"- 考前高清卷：`04-考前测试/original/` 与 `printables/{unit['id']}/04-考前测试/`",
        f"- 延展卷（待出）：`04-考前测试/extended/`",
        f"- 答案：`../../answers/`",
        "",
    ]
    (base / "README.md").write_text("\n".join(lines), encoding="utf-8")


def place_exam_files(unit: dict, base: Path) -> None:
    src_root = ROOT / "_source" / "unit-tests"
    for p in src_root.glob(unit["source_glob"]):
        shutil.copy2(p, base / "04-考前测试" / "source" / p.name)
    # printables already generated under printables/unit-tests — copy into new tree
    out = ROOT / "printables" / unit["id"] / "04-考前测试"
    out.mkdir(parents=True, exist_ok=True)
    legacy = ROOT / "printables" / "unit-tests"
    if legacy.exists():
        for p in legacy.glob(f"{unit['printable_prefix']}*"):
            shutil.copy2(p, out / p.name)
            shutil.copy2(p, base / "04-考前测试" / "original" / p.name)


def save_track(unit: dict, track_folder: str, stem: str, doc: Document) -> None:
    out_dir = ROOT / "printables" / unit["id"] / track_folder
    out_dir.mkdir(parents=True, exist_ok=True)
    docx_path = out_dir / f"{stem}.docx"
    pdf_path = out_dir / f"{stem}.pdf"
    doc.save(str(docx_path))
    print(f"  Wrote {docx_path.relative_to(ROOT)}")
    try:
        docx_to_pdf(docx_path, pdf_path)
        print(f"  Wrote {pdf_path.relative_to(ROOT)}")
    except Exception as e:
        print(f"  PDF skip: {e}", file=sys.stderr)


def main() -> int:
    for unit in UNITS:
        print(f"== {unit['id']} {unit['name']}")
        base = ensure_dirs(unit)
        write_unit_readme(unit, base)
        place_exam_files(unit, base)
        uid = unit["id"]
        save_track(unit, "01-日常计算", f"{uid}-日常计算-15min-模板", build_daily(unit))
        save_track(unit, "02-重难点题型", f"{uid}-重难点题型-15min-模板", build_key(unit))
        save_track(unit, "03-应用题", f"{uid}-应用题-15min-模板", build_app(unit))
        save_track(unit, "05-错题库", f"{uid}-错题重练-模板", build_wrong(unit))
    # tracks overview
    overview = ROOT / "units" / "_五条线说明.md"
    overview.write_text(
        "\n".join(
            [
                "# 数学书桌 · 每单元五条线",
                "",
                "| 线 | 时长 | 怎么用 |",
                "| --- | --- | --- |",
                "| 1 日常计算 | **约 15 分钟** | 提速；可含短口算+少量竖式 |",
                "| 2 重难点题型 | **约 15 分钟** | 对照教材与单元卷易错点专练 |",
                "| 3 应用题 | **约 15 分钟** | 短思维/情境；与计算纸分开 |",
                "| 4 考前测试 | **30–40 分钟** | 上传原卷 → 高清入库 → 再出延展卷 |",
                "| 5 错题库 | 按量 | 计算错 / 概念错分栏订正 |",
                "",
                "题型与重难点一律参考：`units/_textbook/` 教材 PDF + 各单元 `04-考前测试/source/` 试卷图。",
                "",
            ]
        ),
        encoding="utf-8",
    )
    print("done")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
