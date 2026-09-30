# -*- coding: utf-8 -*-
"""生成四年级上 · 2 分钟口算冲刺纸（Word + PDF）+ 参考答案页。天蓝主色。"""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "printables" / "oral"
ANS = ROOT / "answers"
STEM = "oral-g4a-sprint-02min-01"
sys.path.insert(0, str(Path(__file__).resolve().parent))
from theme import BRAND_LINE, RGB_SKY_DEEP, RGB_SKY_DARK  # noqa: E402

# (题面, 答案) — 四上起步适用：表内乘除、两位数加减、整十整百、简单两位数×一位数
PROBLEMS = [
    ("36＋48", 84), ("72－29", 43), ("8×7", 56), ("54÷6", 9),
    ("45＋37", 82), ("90－56", 34), ("6×9", 54), ("63÷7", 9),
    ("28＋47", 75), ("81－35", 46), ("7×8", 56), ("48÷8", 6),
    ("59＋26", 85), ("70－48", 22), ("9×6", 54), ("72÷9", 8),
    ("34＋58", 92), ("93－47", 46), ("5×12", 60), ("80÷5", 16),
    ("120＋80", 200), ("300－150", 150), ("15×4", 60), ("96÷8", 12),
    ("240＋360", 600), ("500－280", 220), ("25×3", 75), ("84÷7", 12),
    ("16×5", 80), ("108÷9", 12), ("45×2", 90), ("200÷4", 50),
    ("18×6", 108), ("360÷6", 60), ("24×4", 96), ("420÷7", 60),
]


def set_run_font(run, name="宋体", size=11, bold=False, color=None):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def add_para(doc, text, *, size=11, bold=False, align=None, space_after=4, color=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(space_after)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return p


def build_sprint() -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(1.6)
    sec.right_margin = Cm(1.6)
    sec.top_margin = Cm(1.3)
    sec.bottom_margin = Cm(1.3)

    add_para(doc, BRAND_LINE, size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
             space_after=2, color=RGB_SKY_DEEP)
    add_para(doc, "四上 · 口算冲刺纸 · 限时 2 分钟", size=14, bold=True,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2, color=RGB_SKY_DARK)
    add_para(
        doc,
        "计时：番茄钟 / 手机倒计时 2:00　|　铃响停笔　|　做完几题写几题，不要赶着乱写",
        size=9, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4, color=RGB_SKY_DEEP,
    )
    add_para(
        doc,
        "姓名：________　日期：________　完成题数：____　正确题数：____　（进步看「2分钟正确题数」）",
        size=10, space_after=8,
    )

    # 4 列 × 9 行
    table = doc.add_table(rows=9, cols=4)
    table.autofit = True
    idx = 0
    for r in range(9):
        for c in range(4):
            cell = table.cell(r, c)
            cell.text = ""
            p = cell.paragraphs[0]
            n = idx + 1
            expr = PROBLEMS[idx][0]
            run = p.add_run(f"{n}. {expr}＝______")
            set_run_font(run, size=11)
            pf = p.paragraph_format
            pf.space_before = Pt(3)
            pf.space_after = Pt(3)
            idx += 1

    add_para(doc, "", size=8, space_after=6)
    add_para(
        doc,
        "订正区（只订正错题，不必整张重做）：________________________________",
        size=10, space_after=2, color=RGB_SKY_DARK,
    )
    add_para(doc, "________________________________________________________________", size=10, space_after=8)
    add_para(
        doc,
        f"{BRAND_LINE}　·　口算与思维题分开练　·　答案见 answers/{STEM}-参考答案.md",
        size=8, align=WD_ALIGN_PARAGRAPH.CENTER, color=RGB_SKY_DEEP,
    )
    return doc


def build_answer_md() -> str:
    lines = [
        f"# 口算冲刺 · 2 分钟 · 参考答案",
        "",
        f"> {BRAND_LINE}　·　`{STEM}`",
        "",
        "批改后只记：**完成题数**、**正确题数**。进步曲线看「2 分钟正确题数」。",
        "",
        "| 题号 | 算式 | 答案 | 题号 | 算式 | 答案 |",
        "| --- | --- | --- | --- | --- | --- |",
    ]
    for i in range(0, 36, 2):
        a_n, a_e, a_a = i + 1, PROBLEMS[i][0], PROBLEMS[i][1]
        b_n, b_e, b_a = i + 2, PROBLEMS[i + 1][0], PROBLEMS[i + 1][1]
        lines.append(f"| {a_n} | {a_e} | {a_a} | {b_n} | {b_e} | {b_a} |")
    lines += ["", "一页打完即可，不必给孩子看太久。", ""]
    return "\n".join(lines)


def docx_to_pdf(docx_path: Path, pdf_path: Path) -> None:
    import win32com.client

    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    try:
        d = word.Documents.Open(str(docx_path.resolve()))
        d.SaveAs(str(pdf_path.resolve()), FileFormat=17)
        d.Close(False)
    finally:
        word.Quit()


def main() -> int:
    OUT.mkdir(parents=True, exist_ok=True)
    ANS.mkdir(parents=True, exist_ok=True)
    docx_path = OUT / f"{STEM}.docx"
    pdf_path = OUT / f"{STEM}.pdf"
    ans_path = ANS / f"{STEM}-参考答案.md"

    build_sprint().save(str(docx_path))
    print(f"Wrote {docx_path}")
    ans_path.write_text(build_answer_md(), encoding="utf-8")
    print(f"Wrote {ans_path}")

    try:
        docx_to_pdf(docx_path, pdf_path)
        print(f"Wrote {pdf_path}")
    except Exception as e:
        print(f"PDF skipped: {e}", file=sys.stderr)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
