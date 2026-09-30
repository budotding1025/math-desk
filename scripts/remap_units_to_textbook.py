# -*- coding: utf-8 -*-
"""按人教版四上教材目录重建单元结构（9 单元）。"""
from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]

BOOK_UNITS = [
    {"id": "u01", "no": 1, "short": "大数的认识", "legacy": "u01"},
    {"id": "u02", "no": 2, "short": "公顷和平方千米", "legacy": None},
    {"id": "u03", "no": 3, "short": "角的度量", "legacy": "u02"},
    {"id": "u04", "no": 4, "short": "三位数乘两位数", "legacy": "u03"},
    {"id": "u05", "no": 5, "short": "平行四边形和梯形", "legacy": "u05"},
    {"id": "u06", "no": 6, "short": "除数是两位数的除法", "legacy": None},
    {"id": "u07", "no": 7, "short": "条形统计图", "legacy": None},
    {"id": "u08", "no": 8, "short": "数学广角——优化", "legacy": None},
    {"id": "u09", "no": 9, "short": "总复习", "legacy": None},
]

CN_NO = ["一", "二", "三", "四", "五", "六", "七", "八", "九"]
TRACKS = ["01-日常计算", "02-重难点题型", "03-应用题", "04-考前测试", "05-错题库"]
FOCUS = {
    "u01": ("读数写数、数位意义", "近似数与改写省略", "情境中的大数；专题「1亿有多大」"),
    "u02": ("公顷/平方千米进率", "面积单位换算与选择", "土地面积估测"),
    "u03": ("角的分类与度量", "内外圈与拼角", "生活中的角"),
    "u04": ("口算与竖式乘法", "积的变化与估算", "乘算应用题"),
    "u05": ("平行与垂直", "平行四边形/梯形判定", "作图与周长"),
    "u06": ("除数两位数的口算竖式", "商的变化与估算", "除法应用题"),
    "u07": ("读条形统计图", "复式条形统计图", "根据数据提出问题"),
    "u08": ("优化思想", "沏茶/烙饼/田忌赛马等", "合理安排时间"),
    "u09": ("数与代数综合", "图形与几何综合", "统计与广角综合"),
}


def set_run_font(run, name="宋体", size=11, bold=False, color=None):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def add_para(doc, text, *, size=11, bold=False, align=None, space_after=6):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=(30, 126, 184) if bold else None)
    return p


def write_stub_exam(unit: dict) -> Path:
    out_dir = ROOT / "printables" / unit["id"] / "04-考前测试"
    out_dir.mkdir(parents=True, exist_ok=True)
    stem = f"U{unit['no']}_四上_第{CN_NO[unit['no']-1]}单元_{unit['short']}_待补"
    docx = out_dir / f"{stem}.docx"
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.6)
    sec.top_margin = sec.bottom_margin = Cm(1.4)
    add_para(doc, "数学书桌 · Math Desk", size=10, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(
        doc,
        f"数学 · 四上 · 第{CN_NO[unit['no']-1]}单元 · {unit['short']}",
        size=16,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    add_para(doc, "时间：30–40 分钟", size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(
        doc,
        "学校____________　四年级____班　姓名____________",
        size=11,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=12,
    )
    add_para(doc, "【说明】本单元考前卷待按教材与原卷补全。以下为练习框架。", size=10, bold=True)
    daily, key, app = FOCUS[unit["id"]]
    add_para(doc, "一、填空 / 计算（对应日常计算）", size=12, bold=True)
    add_para(doc, f"围绕：{daily}")
    add_para(doc, "1. ________________________________")
    add_para(doc, "2. ________________________________")
    add_para(doc, "3. ________________________________", space_after=10)
    add_para(doc, "二、选一选 / 重难点", size=12, bold=True)
    add_para(doc, f"围绕：{key}")
    add_para(doc, "1. （　　）")
    add_para(doc, "2. （　　）", space_after=10)
    add_para(doc, "三、解决问题", size=12, bold=True)
    add_para(doc, f"围绕：{app}")
    add_para(doc, "1. ________________________________")
    add_para(doc, "2. ________________________________")
    doc.save(str(docx))
    return docx


def write_wrong_template(unit: dict) -> Path:
    out_dir = ROOT / "printables" / unit["id"] / "05-错题库"
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"{unit['id']}-错题重练-模板.docx"
    doc = Document()
    add_para(doc, "数学书桌 · 错题重练", size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, f"第{CN_NO[unit['no']-1]}单元 · {unit['short']}", size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, "计算错：", size=12, bold=True)
    add_para(doc, "1. ________________  订正：________________")
    add_para(doc, "概念错：", size=12, bold=True)
    add_para(doc, "1. ________________  订正：________________")
    doc.save(str(path))
    return path


def write_readme(unit: dict):
    unit_dir = ROOT / "units" / f"{unit['id']}-{unit['short']}"
    unit_dir.mkdir(parents=True, exist_ok=True)
    for t in TRACKS:
        (unit_dir / t).mkdir(parents=True, exist_ok=True)
        if t == "04-考前测试":
            (unit_dir / t / "original").mkdir(exist_ok=True)
            (unit_dir / t / "source").mkdir(exist_ok=True)
    daily, key, app = FOCUS[unit["id"]]
    text = f"""# 第{CN_NO[unit['no']-1]}单元 · {unit['short']}

依据：`units/_textbook/` 人教版四年级上册（教育部审定 2022）。

## 重难点
- 日常计算：{daily}
- 重难点：{key}
- 应用：{app}

## 可打印
- `../../printables/{unit['id']}/`
- 考前卷：`04-考前测试/original/` 与 `printables/{unit['id']}/04-考前测试/`
"""
    (unit_dir / "README.md").write_text(text, encoding="utf-8")


def ensure_printables(unit: dict):
    base = ROOT / "printables" / unit["id"]
    for t in TRACKS:
        (base / t).mkdir(parents=True, exist_ok=True)


def move_tree(src: Path, dst: Path):
    if not src.exists():
        return False
    if dst.exists():
        shutil.rmtree(dst)
    dst.parent.mkdir(parents=True, exist_ok=True)
    shutil.move(str(src), str(dst))
    print("moved", src, "->", dst)
    return True


def docx_to_pdf(docx_path: Path, pdf_path: Path, word) -> None:
    d = word.Documents.Open(str(docx_path.resolve()))
    try:
        d.SaveAs(str(pdf_path.resolve()), FileFormat=17)
    finally:
        d.Close(False)


def rename_exam_files(unit: dict, old_no: int | None):
    """把迁入后的 U{old}_ 文件名改成 U{new}_。"""
    if old_no is None or old_no == unit["no"]:
        return
    exam_dir = ROOT / "printables" / unit["id"] / "04-考前测试"
    if not exam_dir.exists():
        return
    old_prefix = f"U{old_no}_"
    new_cn = CN_NO[unit["no"] - 1]
    old_cn = CN_NO[old_no - 1]
    for f in list(exam_dir.iterdir()):
        if not f.is_file():
            continue
        name = f.name
        if name.startswith(old_prefix):
            new_name = name.replace(old_prefix, f"U{unit['no']}_", 1)
            new_name = new_name.replace(f"第{old_cn}单元", f"第{new_cn}单元", 1)
            # 若旧短名与新短名不同也尝试（角/乘等已一致，公顷为新建）
            target = f.with_name(new_name)
            if target != f:
                if target.exists():
                    target.unlink()
                f.rename(target)
                print("renamed", name, "->", new_name)


def rename_wrong_files(unit: dict, old_id: str | None):
    if not old_id or old_id == unit["id"]:
        return
    wrong_dir = ROOT / "printables" / unit["id"] / "05-错题库"
    if not wrong_dir.exists():
        return
    for f in list(wrong_dir.iterdir()):
        if f.is_file() and f.name.startswith(old_id):
            new_name = f.name.replace(old_id, unit["id"], 1)
            target = f.with_name(new_name)
            if target != f:
                if target.exists():
                    target.unlink()
                f.rename(target)
                print("renamed wrong", f.name, "->", new_name)


def sync_original(unit: dict):
    exam_dir = ROOT / "printables" / unit["id"] / "04-考前测试"
    unit_dir = ROOT / "units" / f"{unit['id']}-{unit['short']}"
    orig = unit_dir / "04-考前测试" / "original"
    orig.mkdir(parents=True, exist_ok=True)
    if not exam_dir.exists():
        return
    for f in exam_dir.glob("*"):
        if f.is_file():
            shutil.copy2(f, orig / f.name)


def find_units_tmp(tmp: Path, prefix: str):
    for d in tmp.iterdir():
        if d.is_dir() and d.name.startswith(f"units_{prefix}"):
            return d
    return None


def main():
    tmp = ROOT / "_remap_tmp"
    if tmp.exists():
        shutil.rmtree(tmp)
    tmp.mkdir()

    legacy_print = {
        "u01": "leg_u01",
        "u02": "leg_u02_angles",
        "u03": "leg_u03_mul",
        "u04": "leg_u04_qty",
        "u05": "leg_u05_para",
    }
    for old, tag in legacy_print.items():
        p = ROOT / "printables" / old
        if p.exists():
            move_tree(p, tmp / f"printables_{tag}")

    for d in (ROOT / "units").iterdir():
        if d.is_dir() and d.name.startswith("u0"):
            move_tree(d, tmp / f"units_{d.name}")

    # place printables
    move_tree(tmp / "printables_leg_u01", ROOT / "printables" / "u01")
    ensure_printables({"id": "u02"})
    move_tree(tmp / "printables_leg_u02_angles", ROOT / "printables" / "u03")
    move_tree(tmp / "printables_leg_u03_mul", ROOT / "printables" / "u04")
    move_tree(tmp / "printables_leg_u05_para", ROOT / "printables" / "u05")
    move_tree(tmp / "printables_leg_u04_qty", ROOT / "printables" / "extra-数量关系")
    for uid in ["u06", "u07", "u08", "u09"]:
        ensure_printables({"id": uid})

    # units folders
    place_units = [
        ("u01", "u01"),
        ("u02", None),
        ("u03", "u02"),
        ("u04", "u03"),
        ("u05", "u05"),
        ("u06", None),
        ("u07", None),
        ("u08", None),
        ("u09", None),
    ]
    for new_id, old_id in place_units:
        unit = next(u for u in BOOK_UNITS if u["id"] == new_id)
        dest = ROOT / "units" / f"{new_id}-{unit['short']}"
        if old_id:
            src = find_units_tmp(tmp, old_id)
            if src:
                # rename folder to new chinese name
                move_tree(src, dest)
            else:
                write_readme(unit)
        else:
            write_readme(unit)

    src_qty = find_units_tmp(tmp, "u04")
    if src_qty:
        move_tree(src_qty, ROOT / "units" / "extra-数量关系（非教材单元）")

    # rename migrated exam/wrong filenames
    rename_exam_files(next(u for u in BOOK_UNITS if u["id"] == "u03"), 2)
    rename_exam_files(next(u for u in BOOK_UNITS if u["id"] == "u04"), 3)
    rename_wrong_files(next(u for u in BOOK_UNITS if u["id"] == "u03"), "u02")
    rename_wrong_files(next(u for u in BOOK_UNITS if u["id"] == "u04"), "u03")
    # u05 wrong/exam already U5 / u05

    # stubs + wrong for units missing exam pdf
    import win32com.client

    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    try:
        for unit in BOOK_UNITS:
            ensure_printables(unit)
            write_readme(unit)
            exam_dir = ROOT / "printables" / unit["id"] / "04-考前测试"
            pdfs = list(exam_dir.glob("*.pdf")) if exam_dir.exists() else []
            if not pdfs:
                docx = write_stub_exam(unit)
                pdf = docx.with_suffix(".pdf")
                docx_to_pdf(docx, pdf, word)
                print("stub exam", pdf.name)
            wrong_docx = write_wrong_template(unit)
            wrong_pdf = wrong_docx.with_suffix(".pdf")
            if not wrong_pdf.exists() or wrong_pdf.stat().st_mtime < wrong_docx.stat().st_mtime:
                docx_to_pdf(wrong_docx, wrong_pdf, word)
                print("wrong pdf", wrong_pdf.name)
            sync_original(unit)
    finally:
        word.Quit()

    # note file for demoted 数量关系
    note = ROOT / "printables" / "extra-数量关系" / "README.md"
    note.parent.mkdir(parents=True, exist_ok=True)
    note.write_text(
        "# 数量关系（非教材独立单元）\n\n"
        "人教版四上目录无「数量关系」独立单元；相关素材保留于此，"
        "可作乘法/行程综合补充练习，不进入主路径九单元。\n",
        encoding="utf-8",
    )

    if tmp.exists():
        shutil.rmtree(tmp)
    print("done remap")


if __name__ == "__main__":
    main()
