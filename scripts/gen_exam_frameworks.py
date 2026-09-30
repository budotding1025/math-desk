# -*- coding: utf-8 -*-
"""为 U6–U9 生成清晰可见的 Word/PDF 考前框架（4 页试卷 + 答案），替代近空白的扫描/HTML 框。"""
from __future__ import annotations

import shutil
import sys
from pathlib import Path

import pymupdf
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(Path(__file__).resolve().parent))
from theme import BRAND_LINE, RGB_SKY_DEEP, RGB_SKY_DARK  # noqa: E402

UNITS = {
    "u06": ("第六单元 · 除数是两位数的除法", ["口算与竖式", "试商与调商", "商的变化规律", "除法应用题"]),
    "u07": ("第七单元 · 条形统计图", ["读单式条形图", "复式条形图", "根据数据制图", "从图中提出问题"]),
    "u08": ("第八单元 · 数学广角——优化", ["合理安排时间", "烙饼问题", "对策与游戏", "生活中的优化"]),
    "u09": ("第九单元 · 总复习", ["数与代数", "图形与几何", "统计与概率", "广角与综合"]),
}

STEM = {
    "u06": "U6_四上_第六单元_除数是两位数的除法",
    "u07": "U7_四上_第七单元_条形统计图",
    "u08": "U8_四上_第八单元_数学广角——优化",
    "u09": "U9_四上_第九单元_总复习",
}


def set_run_font(run, name="宋体", size=11, bold=False, color=None):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def add_para(doc, text, *, size=11, bold=False, align=None, space_after=4, space_before=0, color=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return p


def new_doc() -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.6)
    sec.top_margin = sec.bottom_margin = Cm(1.4)
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    return doc


def docx_to_pdf(docx_path: Path, pdf_path: Path, word) -> None:
    d = word.Documents.Open(str(docx_path.resolve()))
    d.SaveAs(str(pdf_path.resolve()), FileFormat=17)
    d.Close(False)


def build_paper(title: str, topics: list[str]) -> Document:
    doc = new_doc()
    for i, topic in enumerate(topics, 1):
        if i > 1:
            doc.add_page_break()
        add_para(doc, BRAND_LINE, size=10, align=WD_ALIGN_PARAGRAPH.CENTER, color=RGB_SKY_DEEP, space_after=2)
        add_para(doc, f"数学 · 四上 · {title}练习", size=15, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, color=RGB_SKY_DARK, space_after=2)
        add_para(
            doc,
            f"时间：30–40 分钟　|　试卷第 {i}/4 页　|　清晰框架（原卷 JPG 待上传后可替换）",
            size=10,
            align=WD_ALIGN_PARAGRAPH.CENTER,
            space_after=6,
        )
        add_para(doc, "学校____________　四年级____班　姓名____________", size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
        add_para(doc, f"本页主题：{topic}", size=12, bold=True, color=RGB_SKY_DARK, space_after=6)
        add_para(doc, "一、填空 / 口算。", size=12, bold=True, color=RGB_SKY_DARK, space_after=4)
        for n in range(1, 7):
            add_para(doc, f"{n}. ________________________________________________")
        add_para(doc, "二、计算 / 作图区。", size=12, bold=True, color=RGB_SKY_DARK, space_before=8, space_after=4)
        add_para(doc, "7. ________________________________________________")
        add_para(doc, "　　（竖式 / 作图请写在下方空白处）", size=10, color=RGB_SKY_DEEP)
        add_para(doc, "")
        add_para(doc, "　" * 40, size=10, space_after=2)
        add_para(doc, "　" * 40, size=10, space_after=2)
        add_para(doc, "　" * 40, size=10, space_after=2)
        add_para(doc, "8. ________________________________________________")
        add_para(doc, "9. ________________________________________________")
        add_para(
            doc,
            f"数学 · 四上 · {title}　第 {i} 页（共 4 页）",
            size=8,
            align=WD_ALIGN_PARAGRAPH.CENTER,
            color=RGB_SKY_DEEP,
            space_before=12,
        )
    return doc


def build_answer(title: str, topics: list[str]) -> Document:
    doc = new_doc()
    add_para(doc, BRAND_LINE, size=10, align=WD_ALIGN_PARAGRAPH.CENTER, color=RGB_SKY_DEEP, space_after=2)
    add_para(doc, f"{title} · 参考答案", size=15, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, color=RGB_SKY_DARK, space_after=4)
    add_para(doc, "单独答案页 · 可只打印这一页", size=10, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
    add_para(doc, "【说明】原卷 JPG 尚未上传；以下为占位，上传后按真题补全。", size=10, color=RGB_SKY_DEEP, space_after=8)
    for i, topic in enumerate(topics, 1):
        add_para(doc, f"第 {i} 页（{topic}）", size=12, bold=True, color=RGB_SKY_DARK, space_before=6)
        add_para(doc, "1–9 题答案：待原卷确定后补全。／可先对照日常变式自批。")
    add_para(doc, "订正记录：________________________________", space_before=12)
    return doc


def main() -> int:
    import win32com.client

    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    try:
        for uid, (title, topics) in UNITS.items():
            d = ROOT / "printables" / uid / "04-考前测试"
            d.mkdir(parents=True, exist_ok=True)
            stem = d / STEM[uid]
            paper_docx = Path(str(stem) + ".docx")
            ans_docx = Path(str(stem) + "_答案.docx")
            paper_pdf = Path(str(stem) + "_clean_paper.pdf")
            ans_pdf = Path(str(stem) + "_clean_answer.pdf")
            print("==", uid, title)
            build_paper(title, topics).save(str(paper_docx))
            build_answer(title, topics).save(str(ans_docx))
            docx_to_pdf(paper_docx, paper_pdf, word)
            docx_to_pdf(ans_docx, ans_pdf, word)

            q_out = Path(str(stem) + "_试卷.pdf")
            a_out = Path(str(stem) + "_答案.pdf")
            full_out = Path(str(stem) + ".pdf")
            shutil.copy2(paper_pdf, q_out)
            shutil.copy2(ans_pdf, a_out)
            full = pymupdf.open()
            qdoc = pymupdf.open(q_out)
            full.insert_pdf(qdoc)
            adoc = pymupdf.open(a_out)
            full.insert_pdf(adoc)
            qdoc.close()
            adoc.close()
            full.save(full_out, deflate=True, garbage=4)
            full.close()
            for p in (paper_pdf, ans_pdf):
                p.unlink(missing_ok=True)
            print("  ok", q_out.name, q_out.stat().st_size)
    finally:
        word.Quit()
    print("done")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
