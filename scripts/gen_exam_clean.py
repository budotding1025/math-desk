# -*- coding: utf-8 -*-
"""考前测试：按 JPG 内容重排为清晰 Word + PDF（非扫描放大）。"""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(Path(__file__).resolve().parent))
from theme import BRAND_LINE, RGB_SKY_DEEP, RGB_SKY_DARK  # noqa: E402

OUT_MAP = {
    "u01": ROOT / "printables" / "u01" / "04-考前测试" / "U1_四上_第一单元_大数的认识",
    "u02": ROOT / "printables" / "u02" / "04-考前测试" / "U2_四上_第二单元_角的度量",
    "u03": ROOT / "printables" / "u03" / "04-考前测试" / "U3_四上_第三单元_三位数乘两位数_部分",
    "u04": ROOT / "printables" / "u04" / "04-考前测试" / "U4_四上_第四单元_数量关系_部分",
    "u05": ROOT / "printables" / "u05" / "04-考前测试" / "U5_四上_第五单元_平行四边形和梯形_部分",
}
DIAG = ROOT / "printables" / "_diagrams"


def add_image(doc, path: Path, width_cm: float = 15.5):
    if not path.exists():
        add_para(doc, f"【缺图：{path.name}】", size=9, color=RGB_SKY_DEEP)
        return
    doc.add_picture(str(path), width=Cm(width_cm))
    p = doc.paragraphs[-1]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)


def set_run_font(run, name="宋体", size=11, bold=False, color=None):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def add_para(doc, text, *, size=11, bold=False, align=None, space_after=4, space_before=0, color=None, first_line=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    if first_line is not None:
        pf.first_line_indent = first_line
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return p


def new_doc() -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.6)
    sec.top_margin = sec.bottom_margin = Cm(1.4)
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    return doc


def paper_header(doc, unit_cn: str, minutes="40"):
    add_para(doc, BRAND_LINE, size=10, align=WD_ALIGN_PARAGRAPH.CENTER, color=RGB_SKY_DEEP, space_after=2)
    add_para(
        doc,
        f"数学 · 四上 · {unit_cn}练习",
        size=16,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        color=RGB_SKY_DARK,
        space_after=2,
    )
    add_para(doc, f"时间：{minutes} 分钟", size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
    add_para(
        doc,
        "学校____________　四年级____班　姓名____________",
        size=11,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=10,
    )


def footer(doc, unit_cn: str, page: int, total: int):
    add_para(
        doc,
        f"数学 · 四上 · {unit_cn}　第 {page} 页（共 {total} 页）　|　清晰重排版（非扫描放大）",
        size=8,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        color=RGB_SKY_DEEP,
        space_after=0,
    )


def section(doc, title: str):
    add_para(doc, title, size=12, bold=True, color=RGB_SKY_DARK, space_after=6, space_before=8)


def page_break(doc):
    doc.add_page_break()


def docx_to_pdf(docx_path: Path, pdf_path: Path, word=None) -> None:
    import win32com.client

    own = word is None
    if own:
        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False
    try:
        d = word.Documents.Open(str(docx_path.resolve()))
        d.SaveAs(str(pdf_path.resolve()), FileFormat=17)
        d.Close(False)
    finally:
        if own:
            word.Quit()


# ——— U1 ———
def build_u01() -> Document:
    doc = new_doc()
    paper_header(doc, "第一单元")
    section(doc, "一、填空。")
    add_para(doc, "1. 10 个一万是（　　），10 个一百万是（　　），一千万是（　　）个十万，一亿是（　　）个一百万。")
    add_para(doc, "2. 在 ○ 里填“＞”“＜”或“＝”。")
    add_para(doc, "　　5430079 ○ 5430790　　1010100 ○ 110100")
    add_para(doc, "　　48 万 ○ 480000　　　　6299000 ○ 630 万")
    add_para(
        doc,
        "3. 马里亚纳海沟最深处为 11034 米。横线上的数读作（　　　　　　　　）米。"
        "它是由（　　）个万和（　　）个一组成的。",
    )
    add_para(
        doc,
        "4. 一个八位数，各个数位上都是 9，再加上 1，得到的数是（　　　　　　）。"
        "新得到的这个数是（　　）位数，最高位在（　　）位上。",
    )
    add_para(
        doc,
        "5. 从月球到地球的平均距离约三十八万四千四百千米。横线上的数写作（　　　　）千米，"
        "省略“万”后面的尾数约是（　　）万千米。",
    )
    add_para(doc, "6. 28058306 这个数中，万级的数字 8 表示（　　　　　　），个级的数字 8 表示（　　　　　　）。")
    add_para(doc, "7. 8□989≈8 万，□里可以填（　　　　）；10□0074000≈11 亿，□里可以填（　　　　）。")
    add_para(
        doc,
        "8. 35□8270060 是一个十位数，□里填（　　）时，这个数最接近 35 亿；"
        "如果这个数最接近 36 亿，□里应填（　　）。",
    )
    add_para(
        doc,
        "9. 算盘是我国古代的伟大发明。如图，算盘表示的数写作______________。"
        "请在下面数线上用“↑”标出这个数的大致位置。",
    )
    add_image(doc, DIAG / "u1_abacus_numline_exact.jpg", 16.0)
    add_para(
        doc,
        "10. 11 颗珠子，放在不同数位上表示的数不同。下图表示五位数 51023。",
    )
    add_image(doc, DIAG / "u1_place_value_exact.jpg", 15.0)
    add_para(doc, "如果用这 11 颗珠子组成一个新的五位数（每个数位上都要有珠子），这个五位数最大是（　　　　）。", space_after=8)

    section(doc, "二、选一选。")
    add_para(doc, "1. 543100 中的“3”表示（　　）。")
    add_para(doc, "　　① 3 个一百　　② 3 个一千　　③ 3 个一万　　④ 3 个十万")
    add_para(doc, "2. 有一个数，万级上是 406，个级上是 273，这个数是（　　）。")
    add_para(doc, "　　① 406273　　② 4060273　　③ 4062730　　④ 40600273")
    add_para(doc, "3. 在数字 4 和 7 之间添（　　）个 0，可以组成四千万零七。")
    add_para(doc, "　　① 4　　② 5　　③ 6　　④ 7")
    footer(doc, "第一单元", 1, 4)

    page_break(doc)
    section(doc, "二、选一选。（续）")
    add_para(doc, "4. 四位同学用不同方法表示 120000，其中不正确的是（　　）。")
    add_image(doc, DIAG / "u1_mc120000_exact.jpg", 15.5)
    add_para(doc, "5. 一个数的近似数是 476 万。如果原来这个数万位上的数字是 6，那么原来千位上的数字最大是（　　）。")
    add_para(doc, "　　① 4　　② 5　　③ 0　　④ 9")
    add_para(doc, "6. 下面四个数中，（　　）可能是图中 M 点表示的数。")
    add_image(doc, DIAG / "u1_numline_M_exact.jpg", 14.0)
    add_para(doc, "　　① 63000　　② 65000　　③ 67000　　④ 69000")

    section(doc, "三、按要求完成下面各题。")
    add_para(doc, "1. 连一连（把读 0 的情况与数连起来）。")
    add_para(doc, "　　一个 0 都不读　　读出一个 0　　读出两个 0")
    add_para(doc, "　　80072050　　80700025　　87025000　　87205000")
    add_para(doc, "2. 填上合适的数。")
    add_para(doc, "　　(1) 379614＝100000×（　）＋（　）×7＋（　）×9＋100×（　）＋10×（　）＋（　）×1")
    add_para(doc, "　　(2) 在数线 50000——60000——70000 上，标出约 55000 与约 63000 的位置（用方框标数）。")
    add_para(doc, "3. 将表格中四大行星到太阳的平均距离按从大到小的顺序排一排。")
    pt = doc.add_table(rows=5, cols=2)
    pt.style = "Table Grid"
    pt.rows[0].cells[0].text = "行星"
    pt.rows[0].cells[1].text = "平均距离（千米）"
    rows = [("地球", "149600000"), ("木星", "778330000"), ("水星", "57910000"), ("金星", "108200000")]
    for i, (a, b) in enumerate(rows, 1):
        pt.rows[i].cells[0].text = a
        pt.rows[i].cells[1].text = b
    add_para(doc, "从大到小：______________ → ______________ → ______________ → ______________", space_after=8)
    footer(doc, "第一单元", 2, 4)

    page_break(doc)
    section(doc, "四、列竖式计算。（带 ※ 的题目要验算）")
    add_para(doc, "1. 76×45＝　　　　　　2. 21.4－7.9＝")
    add_para(doc, "")
    add_para(doc, "3. 572÷7＝　　　　　　※4. 843÷5＝")
    add_para(doc, "（请在下方列竖式，※ 题写出验算）", size=9, color=RGB_SKY_DEEP, space_after=16)

    section(doc, "五、根据要求完成下面各题。")
    add_para(
        doc,
        "2022 年 2 月 4 日，第 24 届冬季奥林匹克运动会在北京开幕。"
        "国家体育场（鸟巢）于 2003 年 12 月 24 日开工建设，2008 年 6 月 28 日竣工，"
        "总投资 2267000000 元，建筑面积 258000 平方米，可容纳约 91000 名观众。"
        "本届冬奥会，中国体育代表团共获得 9 枚金牌、4 枚银牌、2 枚铜牌，共 15 枚奖牌。",
        size=10,
        first_line=Cm(0.74),
        space_after=8,
    )
    add_para(
        doc,
        "1. 2267000000 读作（　　　　　　　　　　　　）。"
        "这个数中，左边的“2”在（　　）位上，表示（　　　　）；"
        "右边的“2”在（　　）位上，表示（　　　　）。"
        "258000 省略万位后面的尾数约是（　　）万。",
    )
    add_para(doc, "2. 选择：① 准确数　　② 近似数")
    add_para(doc, "　　“可容纳约 91000 名观众”中的 91000 是（　　）；“获得 9 枚金牌”中的 9 是（　　）。")
    add_para(
        doc,
        "3. 冬奥会期间，小明听说“带动约 3 亿人参与冰雪运动”。"
        "他想用四舍五入的知识估计：这个“约 3 亿”最多可能表示多少人？"
        "请写出你的想法。",
    )
    add_para(doc, "________________________________________________________________")
    add_para(doc, "________________________________________________________________", space_after=10)
    footer(doc, "第一单元", 3, 4)

    # note: content spans ~3 pages in clean layout; add blank page marker for "共4页" alignment
    page_break(doc)
    add_para(doc, "（本页为作答／竖式续写空白页）", size=10, align=WD_ALIGN_PARAGRAPH.CENTER, color=RGB_SKY_DEEP)
    add_para(doc, "", space_after=200)
    footer(doc, "第一单元", 4, 4)
    return doc


def build_u02() -> Document:
    doc = new_doc()
    paper_header(doc, "第二单元")
    add_para(doc, "【说明】图题已按原卷重绘为高清示意图。", size=9, color=RGB_SKY_DEEP)
    section(doc, "一、填空。")
    add_para(doc, "1. 1 周角＝（　　）平角＝（　　）直角。（　　）角＜（　　）角＜（　　）角＜（　　）角＜周角。")
    add_para(doc, "2. 从早上 6 时到中午 12 时，时针转过（　　）°，相当于 2 个（　　）角的度数。")
    add_para(doc, "3. 下图中箭头所指的位置表示（　　）角；请你用“△”标出周角的位置。")
    add_image(doc, DIAG / "u2_degree_ray_exact.jpg", 15.0)
    add_para(doc, "4. 东东用量角器量角时，错误地把外圈刻度当成内圈刻度，读出的度数是 56°，正确的度数应该是（　　）°。")
    add_para(doc, "5. 如右图，半圆形材料做成的“量角器”被分成 9 个同样大的角。若 ∠1 占其中 2 份，则 ∠1＝（　　）°。")
    add_image(doc, DIAG / "u2_9sector_exact.jpg", 9.0)
    add_para(doc, "6. 古人将太阳周年运动轨迹 360° 划分为 24 等份，每个节气的太阳运动轨迹是（　　）°。")

    section(doc, "二、选一选。")
    add_para(doc, "1. 两个锐角可以组成的角不可能是（　　）。")
    add_para(doc, "　　① 锐角　　② 直角　　③ 钝角　　④ 平角")
    add_para(doc, "2. 如图，收费亭转杆升起过程中，与竖杆形成的角的变化情况为（　　）。")
    add_image(doc, DIAG / "u2_toll_exact.jpg", 14.0)
    add_para(doc, "　　① 直角→钝角→周角　　② 锐角→直角→钝角　　③ 直角→钝角→平角　　④ 锐角→钝角→直角")
    add_para(doc, "3. 椅背与椅面夹角宜在 103°～112°，下面各图中符合要求的是（　　）。")
    add_image(doc, DIAG / "u2_chairs_exact.jpg", 15.0)
    add_para(doc, "4. 长方形纸折叠后，∠1 的度数是（　　）。")
    add_image(doc, DIAG / "u2_fold_rect_exact.jpg", 9.0)
    add_para(doc, "　　① 90°　　② 105°　　③ 135°　　④ 150°")
    add_para(doc, "5. 入门道与地面夹角约 10°，高级道更陡，夹角大约是（　　）。")
    add_image(doc, DIAG / "u2_ski_exact.jpg", 13.0)
    add_para(doc, "　　① 20°　　② 40°　　③ 60°　　④ 80°")
    footer(doc, "第二单元", 1, 4)

    page_break(doc)
    section(doc, "三、解决问题。")
    add_para(doc, "1. 先估计，再量一量图中各角的度数并填写。")
    add_image(doc, DIAG / "u2_measure_angles_exact.jpg", 14.0)
    add_para(doc, "　　∠1（锐角）估计______°，实测______°；∠2（钝角）估计______°，实测______°。")
    add_para(doc, "2. 画出下面的角，并写出是什么角。")
    add_image(doc, DIAG / "u2_draw_boxes_exact.jpg", 14.0)
    add_para(doc, "　　(1) 50°（　　）角　　　　(2) 150°（　　）角")
    add_para(doc, "3. 把一张圆形纸连续对折三次，想一想，填一填。")
    add_image(doc, DIAG / "u2_circle_fold_exact.jpg", 15.5)
    add_para(doc, "4. 填出三角尺上每个角的度数。")
    add_image(doc, DIAG / "u2_set_squares_exact.jpg", 14.0)
    add_para(doc, "　　拼角：按图填写。")
    add_image(doc, DIAG / "u2_compose_angles_exact.jpg", 15.0)
    footer(doc, "第二单元", 2, 4)

    page_break(doc)
    add_para(doc, "5. 求出下面图中指定角的度数。")
    add_image(doc, DIAG / "u2_intersect_exact.jpg", 15.5)
    add_para(doc, "6. 在破损量角器上画一个 75° 的角，并标出度数。（利用两刻度差为 75°）")
    add_image(doc, DIAG / "u2_broken_protractor_exact.jpg", 12.0)
    add_para(doc, "7. 折纸飞机：已知 ∠1＝60°，求 ∠2＝（　　）°。")
    add_image(doc, DIAG / "u2_plane_fold_exact.jpg", 14.0)
    add_para(doc, "8. 放风筝：线长相同。")
    add_image(doc, DIAG / "u2_kites_exact.jpg", 14.0)
    add_para(doc, "　　(1) 量出甲、乙风筝线与地面的夹角：甲（　　）°，乙（　　）°。")
    add_para(doc, "　　(2) 风筝飞的高度和夹角有什么关系？________________________________")
    add_para(doc, "　　(3) 若丙的夹角为 30°，他的风筝比甲、乙高吗？请说明。")
    footer(doc, "第二单元", 3, 4)

    page_break(doc)
    add_para(doc, "（作图与量角空白页：请在本页完成画角、标角）", size=10, align=WD_ALIGN_PARAGRAPH.CENTER, color=RGB_SKY_DEEP)
    add_para(doc, "", space_after=200)
    footer(doc, "第二单元", 4, 4)
    return doc


def build_u03() -> Document:
    doc = new_doc()
    paper_header(doc, "第三单元")
    add_para(doc, "【说明】原卷扫描仅有第 1–2 页；图题（竖式箭头）已补全。", size=9, color=RGB_SKY_DEEP)
    section(doc, "一、填空。")
    add_para(doc, "1. 23 个 12 的和是（　　），17 个 26 的积是（　　）。")
    add_para(doc, "2. 最大的两位数与最小的三位数的积是（　　）。")
    add_para(doc, "3. 457 加上（　　）正好是 82 的 9 倍。")
    add_para(doc, "4. 空间站约 90 分钟绕地球一圈。照这样计算，一天（24 小时）大约能绕地球（　　）圈。")
    add_para(doc, "5. 观察规律填空：")
    add_para(doc, "　　12×9－8＝100")
    add_para(doc, "　　123×9－7＝1100")
    add_para(doc, "　　1234×9－6＝11100")
    add_para(doc, "　　12345×9－5＝（　　）")
    add_para(doc, "　　……")
    add_para(doc, "　　123456789×9－1＝（　　）")
    add_para(doc, "6. 计算 12×14，正确的方法打“✓”。")
    add_para(doc, "　　东东：14＋14＋…＋14（共 12 个）（　）")
    add_para(doc, "　　成成：12×7＋2（　）")
    add_para(doc, "　　思思：14×6＋14×6（　）")
    add_para(doc, "　　乐乐：12×10＋12×4（　）")
    add_para(doc, "7. 一块长方形绿地，长增加 1 米，面积增加 16 平方米；宽增加 1 米，面积增加 25 平方米。这块绿地的面积是（　　）平方米。")
    footer(doc, "第三单元（部分）", 1, 2)

    page_break(doc)
    section(doc, "二、选一选。")
    add_para(doc, "1. 卫星绕地球一圈要 114 分钟，11 圈要多少分钟？竖式 114×11 中，箭头所指部分积表示（　　）。")
    add_image(doc, DIAG / "u3_114x11_exact.jpg", 7.0)
    add_para(doc, "　　① 1 圈 114 分　　② 1 圈 1140 分　　③ 10 圈 114 分　　④ 10 圈 1140 分")
    add_para(doc, "2. 用计算器算 3090×120，误输入为 3090×20。在不清空重输的情况下，接着怎样操作可以得到正确答案？（　　）")
    add_para(doc, "　　① ÷10　　② ×10　　③ ×6　　④ ÷6")
    add_para(doc, "3. 做诗集，每人 32 张纸，298 人大约需要多少张？谁的估算能解决问题？（　　）")
    add_para(doc, "　　① 东东：298×32≈8940（把 32 看作 30）")
    add_para(doc, "　　② 成成：298×32≈8700（298≈290，32≈30）")
    add_para(doc, "　　③ 思思：298×32≈9600（298≈300）")
    add_para(doc, "　　④ 乐乐：298×32≈9000（298≈300，32≈30）")
    footer(doc, "第三单元（部分）", 2, 2)
    return doc


def build_u04() -> Document:
    doc = new_doc()
    paper_header(doc, "第四单元")
    add_para(doc, "【说明】原卷扫描仅有第 1–2 页，本卷按已有内容清晰重排。", size=9, color=RGB_SKY_DEEP)
    section(doc, "一、填空。")
    add_para(doc, "1. 故宫文创店全天接待消费者 1280 人，上午 659 人，下午（　　）人。是根据数量关系（　　）－（　　）＝（　　）来解决问题的。")
    add_para(doc, "2. 图书馆购买了 15 套丛书，每套 108 元，一共花了（　　）元，是根据（　　）×（　　）＝（　　）来解决问题的。")
    add_para(doc, "3. 沪宁高速全长 274 千米，汽车行驶 3 小时后离上海还有 34 千米，平均每小时行驶（　　）千米。")
    add_para(doc, "4. 四年级男生 113 人，比女生少 46 人，四年级共有学生（　　）人。")
    add_para(doc, "5. （1）妈妈说“牛肉真贵”通常指（　　）高。")
    add_para(doc, "　　（2）笔记本 4360 元，买 2 台。已知（　　）和（　　），求（　　）。")
    add_para(doc, "　　（3）校服 120 元，买 5 套。关系式是（　　）。")
    add_para(doc, "6. 汽车 80 千米/时，普通列车 120 千米/时，高铁 230 千米/时，飞机 730 千米/时。")
    add_para(doc, "　　王老师从 A 到 B 共 1260 千米，8:50 出发，14:50 到达，她较可能乘坐（　　）。")
    add_para(doc, "7. 竖式 125×23 中，箭头所指部分积表示（　　）件毛衣的总价，是（　　）元。")
    add_image(doc, DIAG / "u4_125x23_exact.jpg", 7.0)
    footer(doc, "第四单元（部分）", 1, 2)

    page_break(doc)
    section(doc, "二、选一选。")
    add_para(doc, "1. 复兴号约 350 千米/时，15 分钟行驶多少千米？这个问题求的是（　　）。")
    add_para(doc, "　　① 速度　　② 路程　　③ 时间　　④ 车次")
    add_para(doc, "2. 买 10 个篮球共 400 元，每个多少元？求的是（　　）。")
    add_para(doc, "　　① 单价　　② 数量　　③ 总价　　④ 速度")
    add_para(doc, "3. 下面哪个速度最快？（先统一单位再比较）（　　）")
    add_para(doc, "　　① 约 20 m/s　　② 约 5 km/min　　③ 约 900 km/h　　④ 约 90 km/h")
    add_para(doc, "4. 骑行 225 米/分，骑 12 分钟。竖式 225×12 中箭头所指（方框内）一步表示（　　）。")
    add_image(doc, DIAG / "u4_225x12_exact.jpg", 7.0)
    add_para(doc, "　　① 1 分钟路程　　② 2 分钟路程　　③ 10 分钟路程　　④ 12 分钟路程")
    add_para(doc, "5. 下列不能用 15×4 解决的是（　　）。")
    add_para(doc, "　　① 4 条丝带，每条 15 m，总长多少？")
    add_para(doc, "　　② 甲有 15 元，乙是甲的 4 倍，乙有多少？")
    add_para(doc, "　　③ 宽 15 m，长是宽的 4 倍，求面积。")
    add_para(doc, "　　④ 鸡蛋原价 20 元/kg，现价 15 元/kg，买 4 kg 多少钱？")
    footer(doc, "第四单元（部分）", 2, 2)
    return doc


def build_u05() -> Document:
    doc = new_doc()
    paper_header(doc, "第五单元")
    add_para(doc, "【说明】原卷扫描仅有第 1–2 页，本卷按已有内容清晰重排；图题已用高清示意图补全。", size=9, color=RGB_SKY_DEEP)
    section(doc, "一、填空。")
    add_para(doc, "1. 观察下面各图，互相垂直的是（　　），互相平行的是（　　）。")
    add_image(doc, DIAG / "u5_lines_exact.jpg", 16.0)
    # also keep AI-refined version note - prefer exact
    add_para(doc, "2. 在梯形下面的括号里画“○”，在平行四边形下面的括号里画“△”。")
    add_image(doc, DIAG / "u5_shapes_exact.jpg", 16.0)
    add_para(doc, "3. 从直线外一点到这条直线的（　　）线段最短，它的长度叫做点到直线的（　　）。")
    add_para(doc, "4. 只有一组对边平行的四边形叫做（　　）。平行的一组对边分别叫做（　　）和（　　），不平行的两边叫做（　　）。")
    add_para(doc, "5. 梯形最多有（　　）个直角，这样的梯形叫做（　　）梯形。")
    add_para(doc, "6. 等腰梯形的两腰（　　），同一底上的两个底角（　　）。")
    add_para(doc, "7. 平行四边形的邻边分别是 15 cm 和 18 cm，它的周长是（　　）cm。")
    add_para(doc, "8. 如图，写出互相垂直与互相平行的直线。")
    add_image(doc, DIAG / "u5_abcd_hd.jpg", 12.0)
    footer(doc, "第五单元（部分）", 1, 2)

    page_break(doc)
    section(doc, "二、选一选。")
    add_para(doc, "1. 被遮挡图形露出两边平行，它不可能是（　　）。")
    add_image(doc, DIAG / "u5_obscured_hd.jpg", 10.0)
    add_para(doc, "　　① 等腰梯形　　② 平行四边形　　③ 长方形　　④ 直角梯形")
    add_para(doc, "2. 在等腰梯形中画一条直线，不能把它分成两个完全相同的（　　）。")
    add_para(doc, "　　① 梯形　　② 平行四边形　　③ 三角形　　④ 长方形")
    add_para(doc, "3. 两组对边分别平行的四边形是（　　）。")
    add_para(doc, "　　① 直角梯形　　② 平行四边形　　③ 三角形　　④ 等腰梯形")
    add_para(doc, "4. 用圆规比较线段 AB、CD 的长短（开口越大越长），结果是（　　）。")
    add_image(doc, DIAG / "u5_compass_exact.jpg", 13.0)
    add_para(doc, "　　① CD＞AB　　② CD＝AB　　③ CD＜AB　　④ 无法确定")
    add_para(doc, "5. 如图，梯形 ABCD 中 AD∥BC，点 D 沿直线向 A 运动直至重合，图形变化顺序是（　　）。")
    add_image(doc, DIAG / "u5_trap_motion_exact.jpg", 11.0)
    add_para(doc, "　　① 梯形→平行四边形→梯形　　② 梯形→平行四边形→三角形")
    add_para(doc, "　　③ 梯形→三角形→平行四边形→梯形　　④ 梯形→平行四边形→梯形→三角形")

    section(doc, "三、解决问题。")
    add_para(doc, "1. （1）过角内一点 P，分别向两边作垂线。")
    add_image(doc, DIAG / "u5_angle_P_exact.jpg", 10.0)
    add_para(doc, "　　（2）过直线 l 上方点 A、下方点 B，分别作 l 的垂线。这两条垂线的位置关系是（　　）。")
    add_image(doc, DIAG / "u5_line_AB_exact.jpg", 12.0)
    footer(doc, "第五单元（部分）", 2, 2)
    return doc


BUILDERS = {
    "u01": build_u01,
    "u02": build_u02,
    "u03": build_u03,
    "u04": build_u04,
    "u05": build_u05,
}


def main() -> int:
    import win32com.client

    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    try:
        for key, builder in BUILDERS.items():
            stem = OUT_MAP[key]
            stem.parent.mkdir(parents=True, exist_ok=True)
            docx_path = Path(str(stem) + ".docx")
            pdf_path = Path(str(stem) + ".pdf")
            print("==", key, docx_path.name)
            builder().save(str(docx_path))
            # also copy into units/.../original
            unit_dirs = {
                "u01": "u01-大数的认识",
                "u02": "u02-角的度量",
                "u03": "u03-三位数乘两位数",
                "u04": "u04-数量关系",
                "u05": "u05-平行四边形和梯形",
            }
            orig = ROOT / "units" / unit_dirs[key] / "04-考前测试" / "original"
            orig.mkdir(parents=True, exist_ok=True)
            import shutil

            shutil.copy2(docx_path, orig / docx_path.name)
            docx_to_pdf(docx_path, pdf_path, word=word)
            shutil.copy2(pdf_path, orig / pdf_path.name)
            # legacy folder
            legacy = ROOT / "printables" / "unit-tests"
            if legacy.exists():
                shutil.copy2(docx_path, legacy / docx_path.name)
                shutil.copy2(pdf_path, legacy / pdf_path.name)
            print("  ok", pdf_path.name, "size", pdf_path.stat().st_size)
    finally:
        word.Quit()
    print("done")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
